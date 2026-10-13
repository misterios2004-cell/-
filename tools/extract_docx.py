"""Извлекает банк вопросов и учебные материалы из исходного .docx.

Использование:
    python tools/extract_docx.py <путь к .docx>

Результат записывается в attestation/data/bank.json и attestation/data/materials.json.
Требуется пакет python-docx (только для этого скрипта, приложению он не нужен).
"""
import json
import re
import sys
from pathlib import Path

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

OUT_DIR = Path(__file__).resolve().parent.parent / "attestation" / "data"
LETTERS = "абвгде"


def clean(text):
    return re.sub(r"\s+", " ", text.replace(" ", " ")).strip()


def runs(p):
    """Абзац как список сегментов [текст, жирный]; соседние сегменты с одинаковым начертанием склеиваются."""
    segs = []
    for r in p.runs:
        if not r.text:
            continue
        bold = bool(r.bold)
        if segs and segs[-1][1] == bold:
            segs[-1][0] += r.text
        else:
            segs.append([r.text, bold])
    for s in segs:
        s[0] = s[0].replace(" ", " ")
    return segs


def body(document):
    for el in document.element.body.iterchildren():
        tag = el.tag.split("}")[1]
        if tag == "p":
            yield Paragraph(el, document)
        elif tag == "tbl":
            yield Table(el, document)


def table_rows(t):
    rows = []
    for r in t.rows:
        cells = []
        for c in r.cells:
            cells.append("\n".join(clean(p.text) for p in c.paragraphs if clean(p.text)))
        rows.append(cells)
    return rows


def heading_level(p):
    style = p.style.name
    text = clean(p.text)
    if style == "Heading 1" or (style == "Стиль2" and text.startswith("БЛОК")):
        return 1
    if style in ("Heading 2", "Стиль2"):
        return 2
    if style == "Heading 3":
        return 3
    return 0


def parse_closed(t, topic_id, topic_title):
    questions = []
    for row in t.rows[1:]:
        num_cell, q_cell, key_cell = row.cells
        lines = [clean(p.text) for p in q_cell.paragraphs if clean(p.text)]
        text, options = lines[0], []
        for line in lines[1:]:
            m = re.match(r"^([а-е])\)\s*(.*)$", line)
            if not m:
                raise ValueError(f"Не распознан вариант ответа: {line!r}")
            options.append(m.group(2).rstrip(";."))
        key = clean(key_cell.text)
        questions.append({
            "id": f"C{clean(num_cell.text)}",
            "num": int(clean(num_cell.text)),
            "topic": topic_id,
            "topic_title": topic_title,
            "text": text,
            "options": options,
            "answer": LETTERS.index(key),
        })
    return questions


def parse_open(t, set_id):
    items = []
    for row in t.rows[1:]:
        num, q, ref = (clean(c.text) for c in row.cells)
        items.append({"id": f"{set_id}-{num}", "num": int(num), "text": q, "reference": ref})
    return items


def main(path):
    document = docx.Document(path)
    elements = list(body(document))

    # --- банк вопросов (Часть II) ---
    closed_topics, closed = [], []
    open_sets = {}
    part2 = False
    current_h1 = current_h2 = ""
    for el in elements:
        if isinstance(el, Paragraph):
            text = clean(el.text)
            if text == "ЧАСТЬ II":
                part2 = True
            lvl = heading_level(el)
            if lvl == 1:
                current_h1, current_h2 = text, ""
            elif lvl == 2:
                current_h2 = text
            continue
        if not part2:
            continue
        header = [clean(c.text) for c in el.rows[0].cells]
        if header[:3] == ["№", "Вопрос и варианты ответа", "Ключ"]:
            m = re.match(r"^2\.(\d)\.\s*(.*)$", current_h2)
            topic_id = f"1.{m.group(1)}"  # нумерация темы Модуля 1
            closed_topics.append({"id": topic_id, "title": m.group(2)})
            closed.extend(parse_closed(el, topic_id, m.group(2)))
        elif header[:3] == ["№", "Вопрос", "Эталонный ответ"]:
            if current_h1.startswith("3."):
                key = "aspirator"
            elif current_h1.startswith("4."):
                key = "classification"
            else:
                key = {"5.3": "U2-U3", "5.4": "U3-U4", "5.5": "U4-U5"}[current_h2[:3]]
            open_sets[key] = parse_open(el, key)

    bank = {"closed_topics": closed_topics, "closed": closed, "open": open_sets}
    assert len(closed) == 48, len(closed)
    assert [len(open_sets[k]) for k in ("aspirator", "classification", "U2-U3", "U3-U4", "U4-U5")] == [20, 12, 10, 10, 10]

    # --- учебные материалы ---
    chapters = []
    chapter = None
    skip = True
    for el in elements:
        if isinstance(el, Paragraph):
            text = clean(el.text)
            lvl = heading_level(el)
            if text == "0. Как устроена система":
                skip = False
            if lvl == 1 or text == "0. Как устроена система":
                # Пустые формы и служебные разделы в материалы не включаются
                skip_titles = ("Приложение.", "2. Модуль 1", "3. Модуль 2", "4. Проверка", "5. Аттестационные")
                if text.startswith(skip_titles):
                    skip = True
                    continue
                if text.startswith(("1. Порядок применения", "6. Практические")) or not skip:
                    skip = False
                    title = {"1. Порядок применения": "Порядок аттестации",
                             "6. Практические задания по уровням": "Практические задания по уровням"}.get(text, text)
                    chapter = {"title": title, "blocks": []}
                    chapters.append(chapter)
                    continue
            if skip or chapter is None or not text:
                continue
            if text == "Что требуется уточнить до утверждения":
                skip = True
                continue
            if lvl:
                chapter["blocks"].append({"type": "h", "level": lvl, "text": text})
            elif el.style.name == "List Paragraph":
                chapter["blocks"].append({"type": "li", "runs": runs(el)})
            else:
                chapter["blocks"].append({"type": "p", "runs": runs(el)})
        else:
            if skip or chapter is None:
                continue
            chapter["blocks"].append({"type": "table", "rows": table_rows(el)})

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUT_DIR / "bank.json").write_text(json.dumps(bank, ensure_ascii=False, indent=1), encoding="utf-8")
    (OUT_DIR / "materials.json").write_text(json.dumps({"chapters": chapters}, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"Закрытых вопросов: {len(closed)}; открытых наборов: {len(open_sets)}; глав материалов: {len(chapters)}")


if __name__ == "__main__":
    main(sys.argv[1])
